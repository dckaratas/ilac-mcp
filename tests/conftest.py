"""Gerçek TİTCK dosyalarının yapısını taklit eden küçük örnek XLSX'ler.

Gerçek listeler repoya konmaz; başlık metinleri ve sayfa adları 23.10.2026
ve 20.10.2026 tarihli dosyalardan birebir alınmıştır.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

ASKI_HEADER = (
    "RUHSATI ASKIDA OLMAYAN ÜRÜN: 0\r\nMADDE-23 GEREKÇELİ ASKIDA OLAN ÜRÜN: 1\r\n"
    "FARMAKOVİJİLANS GEREKÇELİ ASKIDA OLAN ÜRÜN: 2\r\nMADDE-22 GEREKÇELİ ASKIDA OLAN ÜRÜN: 3"
)


@pytest.fixture
def ruhsatli_xlsx(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "RUHSATLI ÜRÜNLER LİSTESİ"
    ws.append(["RUHSATLI BEŞERİ TIBBİ ÜRÜNLER LİSTESİ", "İRD-LST-15 /31.12.2019/Rev.01/17.11.2023"])
    ws.append([
        "SIRA NO", "BARKOD", "ÜRÜN ADI", "ETKİN MADDE", "ATC KODU", "RUHSAT SAHİBİ", "RUHSAT TARİHİ",
        "RUHSAT NUMARASI", "DEĞİŞİKLİK", "Bu hafta değişiklik yapılan ürünler 1 ile belirtilmiştir. ",
        "DEĞİŞİKLİK TARİHİ", ASKI_HEADER, "ASKIYA ALINMA TARİHİ",
    ])
    ws.append([1, "8699508010752", "ONADRON 0.75 MG TABLET, 100 TABLET", "dekzametazon", "H02AB02",
               "MENARİNİ SAĞLIK VE İLAÇ SAN. TİC. A.Ş.", 22190, "55/44", None, None, 43182, 0, None])
    ws.append([2, "8699508010134", "ONADRON 0.75 MG TABLET, 20 TABLET", "dekzametazon", "H02AB02",
               "MENARİNİ SAĞLIK VE İLAÇ SAN. TİC. A.Ş.", 22190, "55/44", None, None, None, 0, None])
    ws.append([3, "8699000000011", "DEKORT 0.75 MG TABLET, 20 TABLET", "dekzametazon", "H02AB02 ",
               "ÖRNEK İLAÇ A.Ş.", 40000, "100/1", None, None, None, 2, 46000])
    ws.append([4, "8699000000028", "DEKORT 4 MG/ML ENJEKSİYONLUK ÇÖZELTİ", "dekzametazon", "H02AB02",
               "ÖRNEK İLAÇ A.Ş.", 40000, "100/2", None, None, None, 0, None])
    ws.append([5, "8699000000035", "VERMİDON 500/30 MG TABLET, 20 TABLET", "parasetamol, kafein", "N02BE51",
               "SANDOZ İLAÇ SANAYİ VE TİCARET A.Ş.", 30000, "200/1", None, None, None, 0, None])
    path = tmp_path / "ruhsatli.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def erecete_xlsx(tmp_path: Path) -> Path:
    wb = Workbook()
    head = ["İlaç Adı", "Barkod", "ATC Kodu", "ATC Adı", "Firma Adı", "Reçete Türü", "Durumu", "Açıklama"]
    aktif = wb.active
    aktif.title = "AKTİF ÜRÜNLER LİSTESİ"
    aktif.append(["SKRS E-REÇETE İLAÇ VE DİĞER FARMASÖTİK ÜRÜNLER LİSTESİ", " AİK-LST-02/A/ 11.11.2019/ Rev. 00"])
    aktif.append(["AKTİF ÜRÜNLER LİSTESİ (13.10.2026-19.10.2026 arası değişiklikler işlenmiştir.)"])
    aktif.append(head + ["Temel İlaç Listesi Durumu", "Çocuk Temel İlaç\r\nListesi Durumu",
                         "Yenidoğan\r\nTemel İlaç\r\nListesi Durumu", "Aktif Ürünler Listesine\r\nAlındığı Tarih"])
    aktif.append(["ONADRON 0.75 MG 20 TABLET", 8699508010134, "H02AB02", "dexamethasone", "MENARİNİ", "Normal",
                  "Aktif", None, "1", "0", "0", 46315])
    aktif.append(["DEKORT 0.75 MG 20 TABLET", 8699000000011, "H02AB02", "dexamethasone", "ÖRNEK", "Normal",
                  "Aktif", None, "0", "0", "0", 46315])
    aktif.append(["DEKORT 4 MG/ML AMPUL", 8699000000028, "H02AB02", "dexamethasone", "ÖRNEK", "Normal",
                  "Aktif", None, "0", "0", "0", 46315])
    aktif.append(["SADECE ERECETE 10 MG TABLET", 8699000000042, "C10AA05", "atorvastatin", "ÖRNEK", "Normal",
                  "Aktif", None, "0", "0", "0", 46315])
    pasif = wb.create_sheet("PASİF ÜRÜNLER LİSTESİ")
    pasif.append(["SKRS E-REÇETE İLAÇ VE DİĞER FARMASÖTİK ÜRÜNLER LİSTESİ"])
    pasif.append(["PASİF ÜRÜNLER LİSTESİ"])
    pasif.append(head + ["Pasif Ürünler Listesine\r\nAlındığı Tarih"])
    pasif.append(["ONADRON 0.75 MG 100 TABLET", 8699508010752, "H02AB02", "dexamethasone", "MENARİNİ", "Normal",
                  "Pasif", None, 46308])
    alinacak = wb.create_sheet("PASİFE ALINACAK ÜRÜNLER")
    alinacak.append(["SKRS"])
    alinacak.append(["PASİFE ALINACAK ÜRÜNLER LİSTESİ"])
    alinacak.append(head)
    yeni = wb.create_sheet("LİSTEYE YENİ EKLENEN ÜRÜNLER")
    yeni.append(["SKRS"])
    yeni.append(["LİSTEYE YENİ EKLENEN ÜRÜNLER LİSTESİ"])
    yeni.append(head)
    yeni.append(["SADECE ERECETE 10 MG TABLET", 8699000000042, "C10AA05", "atorvastatin", "ÖRNEK", "Normal",
                 "Aktif", "Ekleme yeni ürün/Aktif ürünler listesine eklendi."])
    path = tmp_path / "erecete.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def conn(tmp_path: Path, ruhsatli_xlsx: Path, erecete_xlsx: Path):
    from ilac_mcp import db
    from ilac_mcp.sync import build_from_files

    target = tmp_path / "ilac.db"
    build_from_files(ruhsatli_xlsx, erecete_xlsx, target, {"ruhsatli_liste_tarihi": "2026-10-23"})
    c = db.connect(target)
    yield c
    c.close()
