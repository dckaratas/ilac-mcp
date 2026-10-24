"""TİTCK XLSX listelerini kayıtlara dönüştürme.

Sütunlar başlık metniyle (katlanmış hâliyle) eşlenir; sıra değişse de çalışır.
"""

from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .normalize import clean, fold, to_iso_date

Record = dict[str, Any]

# katlanmış başlık önekiyle eşleşen alan adları
RUHSATLI_COLUMNS = {
    "barkod": "barkod",
    "urun adi": "ad",
    "etkin madde": "etkin_madde",
    "atc kodu": "atc",
    "ruhsat sahibi": "ruhsat_sahibi",
    "ruhsat tarihi": "ruhsat_tarihi",
    "ruhsat numarasi": "ruhsat_no",
    "degisiklik tarihi": "degisiklik_tarihi",
    "ruhsati askida olmayan": "askida_durum",
    "askiya alinma tarihi": "askiya_alinma_tarihi",
}

ERECETE_COLUMNS = {
    "ilac adi": "ad",
    "barkod": "barkod",
    "atc kodu": "atc",
    "atc adi": "atc_adi",
    "firma adi": "firma",
    "recete turu": "recete_turu",
    "durumu": "durum",
    "aciklama": "aciklama",
    "temel ilac listesi": "temel_ilac",
    "cocuk temel ilac": "cocuk_temel_ilac",
    "yenidogan temel ilac": "yenidogan_temel_ilac",
    "aktif urunler listesine": "liste_tarihi",
    "pasif urunler listesine": "liste_tarihi",
}

ASKI_DURUMLARI = {
    0: None,
    1: "Madde-23 gerekçeli askıda",
    2: "Farmakovijilans gerekçeli askıda",
    3: "Madde-22 gerekçeli askıda",
}

DATE_FIELDS = {"ruhsat_tarihi", "degisiklik_tarihi", "askiya_alinma_tarihi", "liste_tarihi"}
FLAG_FIELDS = {"temel_ilac", "cocuk_temel_ilac", "yenidogan_temel_ilac"}


def _match_columns(header: tuple[Any, ...], mapping: dict[str, str]) -> dict[int, str]:
    cols: dict[int, str] = {}
    for idx, cell in enumerate(header):
        key = fold(clean(cell))
        for prefix, field in mapping.items():
            if key.startswith(prefix) and field not in cols.values():
                cols[idx] = field
                break
    return cols


def _iter_sheet(ws, mapping: dict[str, str], scan_rows: int = 10) -> Iterator[Record]:
    rows = ws.iter_rows(values_only=True)
    cols: dict[int, str] = {}
    for _ in range(scan_rows):
        header = next(rows, None)
        if header is None:
            return
        cols = _match_columns(header, mapping)
        if "barkod" in cols.values() and "ad" in cols.values():
            break
    else:
        return
    for row in rows:
        rec: Record = {}
        for idx, field in cols.items():
            raw = row[idx] if idx < len(row) else None
            if field in DATE_FIELDS:
                rec[field] = to_iso_date(raw)
            elif field in FLAG_FIELDS:
                rec[field] = clean(raw) == "1"
            else:
                rec[field] = clean(raw)
        if rec.get("barkod") and rec.get("ad"):
            rec["barkod"] = rec["barkod"].split(".")[0]
            yield rec


def _sheet(wb, *needles: str, exclude: str | None = None):
    for ws in wb.worksheets:
        name = fold(ws.title)
        if all(n in name for n in needles) and not (exclude and exclude in name):
            return ws
    return None


def parse_ruhsatli(path: Path) -> list[Record]:
    """Ruhsatlı Beşeri Tıbbi Ürünler Listesi'nin ana sayfasını okur."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = _sheet(wb, "ruhsatli urunler") or wb.worksheets[0]
        out = []
        for rec in _iter_sheet(ws, RUHSATLI_COLUMNS):
            code = rec.get("askida_durum")
            rec["askida_durum"] = ASKI_DURUMLARI.get(int(code), f"Askıda (kod {code})") if code and code.isdigit() else None
            if rec.get("atc"):
                rec["atc"] = rec["atc"].upper()
            out.append(rec)
        return out
    finally:
        wb.close()


def parse_erecete(path: Path) -> dict[str, list[Record]]:
    """SKRS E-Reçete listesini aktif, pasif ve haftalık değişiklik kayıtlarına ayırır."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = {
            "aktif": _sheet(wb, "aktif urunler"),
            "pasif": _sheet(wb, "pasif urunler", exclude="alinacak"),
            "yeni": _sheet(wb, "yeni eklenen"),
            "degisen": _sheet(wb, "degisiklik yapilan"),
        }
        return {key: list(_iter_sheet(ws, ERECETE_COLUMNS)) if ws else [] for key, ws in sheets.items()}
    finally:
        wb.close()
